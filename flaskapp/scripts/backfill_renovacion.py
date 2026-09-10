# flaskapp/scripts/backfill_renovacion.py
"""
Script de UNA SOLA VEZ para corregir el historial de renovaciones de
pólizas ya existentes, y para dejar documentados (sin tocar nada) los
demás tipos de inconsistencia que se fueron encontrando al revisar los
datos reales.

Contexto: cuando se renueva una póliza, el código nunca guardó en la
póliza VIEJA el folio de la póliza NUEVA que la reemplazó (el campo
`renovacion`) -- solo marcaba `Poliza_renovada = 'Si'`, sin decir a
cuál. Esto ya se corrigió hacia adelante en `app/polizas/routes.py`
(función create()), pero las renovaciones pasadas quedaron con datos
incompletos o inconsistentes de varias formas distintas.

Se identificaron 7 tipos de caso al revisar los datos reales:

  CASO 1 (SE CORRIGE): Poliza_renovada='Si' pero `renovacion` vacío --
      se busca la póliza nueva (la que tiene poliza_anterior == mi
      folio) y se llena `renovacion` con su folio.

  CASO 2 (SE OMITE POR COMPLETO): poliza_anterior lleno pero esa póliza
      no existe en el sistema. Confirmado con el dueño del proyecto:
      son renovaciones de pólizas que nunca se importaron al sistema
      nuevo (existieron en un sistema anterior) -- no es un error, no
      hay nada que corregir, así que ni siquiera se reporta.

  CASO 3 (SE CORRIGE -- el más importante): poliza_anterior lleno, la
      póliza vieja SÍ existe en el sistema, pero esa vieja tiene
      Poliza_renovada != 'Si'. Es el caso más importante porque el
      sistema "no sabe" que esa póliza ya fue renovada -- alguien
      podría intentar renovarla otra vez sin que salte la validación
      de renovación duplicada. Se corrige poniendo Poliza_renovada='Si'
      y renovacion=<folio de la nueva>.

  CASO 4 (SOLO SE REPORTA): el campo `renovacion` sí tiene un folio
      capturado, pero ese folio no existe en el sistema. Puede ser un
      folio mal escrito, o corrupción de datos (ver Caso 6) -- se
      reporta para buscar el número correcto y actualizarlo a mano.

  CASO 5 (SOLO SE REPORTA): la vieja dice "me renovó la póliza X", pero
      esa póliza X no confirma la relación (su poliza_anterior no
      apunta de regreso a la vieja, o apunta a otro folio). Se
      encontraron dos sub-patrones al revisar los datos reales:
        a) Corrupción de Excel: el folio en X.poliza_anterior quedó en
           notación científica (ej. '2.88E+12') en vez del número real.
        b) Reemisión de póliza a mitad de vigencia: se confirmó con un
           caso real (folios 18602508 -> 18602510 -> 18602511) que a
           veces una póliza se reemite con folio nuevo manteniendo el
           mismo período de vigencia (no es una renovación por tiempo),
           y esos casos son demasiado particulares para resolverlos
           automático -- necesitan revisión uno por uno.

  CASO 6 (SOLO SE REPORTA, separado en dos grupos): el mismo folio
      aparece en el campo poliza_anterior de 2 o más pólizas nuevas
      distintas. Se separa en:
        a) Probable corrupción de Excel (el folio duplicado tiene
           forma de notación científica, ej. '1.50E+11') -- muy
           probablemente son pólizas distintas que se corrompieron al
           mismo texto por accidente, no duplicados reales.
        b) Folio con formato normal -- sí podría ser un error real
           (dos pólizas nuevas reclamando la misma póliza vieja), se
           reporta para revisar caso por caso.

  CASO 7 (SOLO SE REPORTA): la póliza tiene TANTO poliza_anterior como
      renovacion llenos, pero NINGUNO de los dos folios existe en el
      sistema (ni siquiera con una búsqueda flexible que ignora
      espacios/guiones/mayúsculas) -- está "atrapada en medio" de un
      historial que empieza y termina fuera de este sistema.

Genera un archivo .xlsx (Excel) con una hoja por cada caso, para poder
revisar cada uno por separado sin que se mezclen.

USO (parado en la carpeta flaskapp/):

    Modo de prueba, no guarda nada, solo genera el Excel:
        python scripts/backfill_renovacion.py

    Modo real, aplica los casos 1 y 3 en la base de datos (los demás
    casos NUNCA se tocan automáticamente, solo se reportan):
        python scripts/backfill_renovacion.py --aplicar
"""
import sys
import os
import re
from datetime import datetime
from collections import defaultdict

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import app, db
from app.models import Poliza

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
except ImportError:
    print("Falta la librería openpyxl. Instálala con:")
    print("    pip install openpyxl==3.1.5")
    sys.exit(1)


RE_NOTACION_CIENTIFICA = re.compile(r'^\d\.\d+E\+\d+$', re.IGNORECASE)


def _normalizar(folio):
    """Para la búsqueda flexible del Caso 7: mayúsculas, sin espacios/guiones."""
    if not folio:
        return None
    return re.sub(r'[\s\-_./]', '', folio.upper().strip())


def backfill(aplicar=False):
    with app.app_context():
        todas = Poliza.query.all()

        por_folio = defaultdict(list)
        por_folio_normalizado = defaultdict(list)
        for p in todas:
            por_folio[p.poliza].append(p)
            por_folio_normalizado[_normalizar(p.poliza)].append(p)

        def existe(folio):
            return bool(folio) and folio in por_folio

        def buscar_nueva_de(vieja_folio):
            candidatas = [p for p in todas if p.poliza_anterior == vieja_folio]
            return candidatas[0] if candidatas else None

        # ---- CASO 1: Poliza_renovada='Si' pero renovacion vacío ----
        caso1 = []
        for p in todas:
            if p.Poliza_renovada == 'Si' and not p.renovacion:
                nueva = buscar_nueva_de(p.poliza)
                caso1.append({'poliza': p, 'nueva_encontrada': nueva})

        # ---- CASO 3: poliza_anterior existe, pero la vieja no sabía que la renovaron ----
        caso3 = []
        for p in todas:
            if not p.poliza_anterior:
                continue
            vieja_lista = por_folio.get(p.poliza_anterior)
            if not vieja_lista:
                continue  # Caso 2, se omite aquí
            for vieja in vieja_lista:
                if vieja.Poliza_renovada != 'Si':
                    caso3.append({'nueva': p, 'vieja': vieja})

        # ---- CASO 4: renovacion apunta a un folio que no existe ----
        caso4 = [p for p in todas if p.renovacion and not existe(p.renovacion)]

        # ---- CASO 5: el link no es de ida y vuelta ----
        caso5 = []
        for p in todas:
            if p.Poliza_renovada == 'Si' and p.renovacion and existe(p.renovacion):
                for nueva in por_folio[p.renovacion]:
                    if nueva.poliza_anterior != p.poliza:
                        es_cientifica = bool(
                            nueva.poliza_anterior and RE_NOTACION_CIENTIFICA.match(nueva.poliza_anterior))
                        caso5.append({
                            'vieja': p, 'nueva': nueva,
                            'probable_corrupcion_excel': es_cientifica,
                        })

        # ---- CASO 6: mismo poliza_anterior reclamado 2+ veces ----
        reclamos = defaultdict(list)
        for p in todas:
            if p.poliza_anterior:
                reclamos[p.poliza_anterior].append(p)
        caso6 = {folio: lista for folio, lista in reclamos.items() if len(lista) > 1}
        caso6_excel = {f: l for f, l in caso6.items() if RE_NOTACION_CIENTIFICA.match(f)}
        caso6_normal = {f: l for f, l in caso6.items() if f not in caso6_excel}

        # ---- CASO 7: atrapada en medio ----
        caso7 = []
        for p in todas:
            if p.poliza_anterior and p.renovacion:
                if not existe(p.poliza_anterior) and not existe(p.renovacion):
                    ant_norm = bool(por_folio_normalizado.get(_normalizar(p.poliza_anterior)))
                    ren_norm = bool(por_folio_normalizado.get(_normalizar(p.renovacion)))
                    if not ant_norm and not ren_norm:
                        caso7.append(p)

        # ================== Aplicar Caso 1 y Caso 3 ==================
        aplicadas_caso1 = 0
        aplicadas_caso3 = 0
        if aplicar:
            for item in caso1:
                if item['nueva_encontrada']:
                    item['poliza'].renovacion = item['nueva_encontrada'].poliza
                    aplicadas_caso1 += 1
            for item in caso3:
                item['vieja'].Poliza_renovada = 'Si'
                item['vieja'].renovacion = item['nueva'].poliza
                aplicadas_caso3 += 1
            db.session.commit()

        # ================== Resumen en consola ==================
        print("=" * 70)
        print(f"CASO 1 (Poliza_renovada=Si, renovacion vacío): {len(caso1)}")
        print(f"  Con póliza nueva encontrada (se puede corregir): "
              f"{sum(1 for i in caso1 if i['nueva_encontrada'])}")
        print("CASO 2 (poliza_anterior no existe): SE OMITE (confirmado: historial externo)")
        print(f"CASO 3 (anterior existe, no sabía que la renovaron): {len(caso3)}")
        print(f"CASO 4 (renovacion apunta a folio inexistente): {len(caso4)}")
        print(f"CASO 5 (link no es de ida y vuelta): {len(caso5)}")
        print(f"  De esos, probable corrupción de Excel: "
              f"{sum(1 for i in caso5 if i['probable_corrupcion_excel'])}")
        print(f"CASO 6 (mismo poliza_anterior reclamado 2+ veces): {len(caso6)} folios")
        print(f"  Probable corrupción de Excel: {len(caso6_excel)}")
        print(f"  Formato normal (revisar a mano): {len(caso6_normal)}")
        print(f"CASO 7 (atrapada en medio, ningún extremo existe): {len(caso7)}")
        print("=" * 70)

        if aplicar:
            print(f"\n✅ Se corrigieron {aplicadas_caso1} póliza(s) del Caso 1 "
                  f"y {aplicadas_caso3} del Caso 3. Cambios guardados.")
        else:
            print("\n⚠️  Esto fue un DRY-RUN — no se guardó nada todavía.")
            print("    Si el reporte se ve bien, vuelve a correr con --aplicar")

        # ================== Excel con una hoja por caso ==================
        wb = Workbook()
        ws_resumen = wb.active
        ws_resumen.title = "Resumen"
        encabezado_fill = PatternFill(start_color="C94A1C", end_color="C94A1C", fill_type="solid")
        encabezado_font = Font(color="FFFFFF", bold=True)

        def nueva_hoja(nombre):
            return wb.create_sheet(nombre)

        def escribir_encabezados(ws, columnas):
            ws.append(columnas)
            for cell in ws[1]:
                cell.fill = encabezado_fill
                cell.font = encabezado_font

        escribir_encabezados(ws_resumen, ["Caso", "Descripción", "Cantidad", "Acción"])
        ws_resumen.append(["1", "Renovada pero no dice con cuál", len(caso1),
                            f"Corregido ({aplicadas_caso1})" if aplicar else "Se corregiría"])
        ws_resumen.append(["2", "Póliza anterior no existe (historial externo)",
                            "—", "Se omite, no es un error"])
        ws_resumen.append(["3", "Anterior existe, no sabía que la renovaron", len(caso3),
                            f"Corregido ({aplicadas_caso3})" if aplicar else "Se corregiría"])
        ws_resumen.append(["4", "Renovación apunta a folio inexistente", len(caso4), "Solo reporte"])
        ws_resumen.append(["5", "El link no es de ida y vuelta", len(caso5), "Solo reporte"])
        ws_resumen.append(["6", "Mismo folio anterior reclamado 2+ veces", len(caso6), "Solo reporte"])
        ws_resumen.append(["7", "Atrapada en medio (ningún extremo existe)", len(caso7), "Solo reporte"])
        ws_resumen.column_dimensions['B'].width = 45

        ws1 = nueva_hoja("1 - Renovada sin decir cual")
        escribir_encabezados(ws1, ["Póliza (vieja) id", "Folio", "Póliza nueva encontrada",
                                    "¿Se corrigió?"])
        for item in caso1:
            p = item['poliza']
            nueva = item['nueva_encontrada']
            ws1.append([p.id, p.poliza, nueva.poliza if nueva else "(no se encontró)",
                        "Sí" if (aplicar and nueva) else ("Se corregiría" if nueva else "No se pudo")])

        ws3 = nueva_hoja("3 - Anterior no sabia")
        escribir_encabezados(ws3, ["Póliza vieja id", "Folio vieja", "Poliza_renovada (antes)",
                                    "Póliza nueva", "Folio nueva"])
        for item in caso3:
            vieja, nueva = item['vieja'], item['nueva']
            ws3.append([vieja.id, vieja.poliza, vieja.Poliza_renovada, nueva.id, nueva.poliza])

        ws4 = nueva_hoja("4 - Renovacion inexistente")
        escribir_encabezados(ws4, ["Póliza id", "Folio", "Renovación (folio buscado, no existe)"])
        for p in caso4:
            ws4.append([p.id, p.poliza, p.renovacion])

        ws5 = nueva_hoja("5 - Link no bidireccional")
        escribir_encabezados(ws5, ["Vieja id", "Vieja folio", "Vieja dice que la renovó",
                                    "Esa póliza (id)", "poliza_anterior de esa póliza",
                                    "Probable corrupción Excel"])
        for item in caso5:
            vieja, nueva = item['vieja'], item['nueva']
            ws5.append([vieja.id, vieja.poliza, vieja.renovacion, nueva.id,
                        nueva.poliza_anterior or "(vacío)",
                        "Sí" if item['probable_corrupcion_excel'] else "No -- revisar a mano"])

        ws6 = nueva_hoja("6 - Folio duplicado")
        escribir_encabezados(ws6, ["Folio anterior reclamado", "Cuántas veces",
                                    "Pólizas que lo reclaman (id:folio)", "Tipo"])
        for folio, lista in caso6_excel.items():
            detalle = "; ".join(f"{p.id}:{p.poliza}" for p in lista)
            ws6.append([folio, len(lista), detalle, "Probable corrupción Excel"])
        for folio, lista in caso6_normal.items():
            detalle = "; ".join(f"{p.id}:{p.poliza}" for p in lista)
            ws6.append([folio, len(lista), detalle, "Revisar a mano"])

        ws7 = nueva_hoja("7 - Atrapada en medio")
        escribir_encabezados(ws7, ["Póliza id", "Folio", "poliza_anterior (no existe)",
                                    "renovacion (no existe)", "Poliza_renovada"])
        for p in caso7:
            ws7.append([p.id, p.poliza, p.poliza_anterior, p.renovacion, p.Poliza_renovada])

        for ws in wb.worksheets:
            for col in ws.columns:
                max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 60)

        nombre_archivo = f"reporte_backfill_renovacion_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        wb.save(nombre_archivo)
        print(f"\n📄 Reporte Excel guardado en: {nombre_archivo}")
        print("    (una hoja por caso, más un resumen)")


if __name__ == '__main__':
    modo_aplicar = '--aplicar' in sys.argv
    backfill(aplicar=modo_aplicar)
